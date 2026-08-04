from fastapi import FastAPI
from fastapi.responses import StreamingResponse, HTMLResponse
from pydantic import BaseModel
from datetime import datetime
import openpyxl
from openpyxl.styles import Alignment
import io

app = FastAPI()

class DataWarga(BaseModel):
    nama: str
    nik: str
    no_kk: str
    no_hp: str
    jenis_permohonan: str
    persyaratan: str = ""

@app.get("/", response_class=HTMLResponse)
async def baca_halaman_web():
    with open("index.html", "r", encoding="utf-8") as f:
        return f.read()

@app.post("/generate-f102")
async def generate_excel(data: DataWarga):
    nama_file_template = "F-1.02 - Formulir Pendaftaran Peristiwa Kependudukan.xlsx"
    wb = openpyxl.load_workbook(nama_file_template)
    ws = wb.active

    # 1. Mengisi Data Diri
    ws['F5'] = data.nama
    ws['F6'] = data.nik
    ws['F7'] = data.no_kk
    ws['F8'] = data.no_hp

    # 2. Logika Centang Jenis Permohonan
    if data.jenis_permohonan:
        sel_centang = data.jenis_permohonan.split(",")
        for sel in sel_centang:
            if sel.strip():
                ws[sel.strip()] = '✔'

    # 3. Logika Centang Persyaratan Lampiran
    if data.persyaratan:
        sel_syarat = data.persyaratan.split(",")
        for sel in sel_syarat:
            if sel.strip():
                ws[sel.strip()] = '✔'

    # 4. Format Tanggal Otomatis (Rata Tengah)
    bulan_indo = ["", "Januari", "Februari", "Maret", "April", "Mei", "Juni", 
                  "Juli", "Agustus", "September", "Oktober", "November", "Desember"]
    tgl_sekarang = datetime.now()
    teks_tanggal = f"Jambi, {tgl_sekarang.day} {bulan_indo[tgl_sekarang.month]} {tgl_sekarang.year}"
    
    ws['K41'] = None
    ws['L41'] = teks_tanggal
    ws['L41'].alignment = Alignment(horizontal='center')

    # 5. Menaruh Nama Pemohon (Rata Tengah)
    ws['D46'] = data.nama
    ws['D46'].alignment = Alignment(horizontal='center')

    # Simpan dan kirim ke browser
    virtual_file = io.BytesIO()
    wb.save(virtual_file)
    virtual_file.seek(0)

    headers = {
        'Content-Disposition': f'attachment; filename="F102_{data.nama}.xlsx"'
    }
    return StreamingResponse(
        virtual_file, 
        headers=headers, 
        media_type='application/vnd.openxmlformats-officedocument.spreadsheetml.sheet'
    )